"""Excel writer -- produces a workbook laid out like the vendor's dashboards."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="0B6B57")
HEAD_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=9)
TOTAL_FONT = Font(name="Arial", size=9, bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="E3EFEA")
THIN = Side(style="thin", color="DDE2DD")
BORDER = Border(bottom=THIN)

MONEY_FMT = '#,##0.00;[Red](#,##0.00);-'
INT_FMT = '#,##0;[Red](#,##0);-'
PCT_FMT = '0.00%'
PCT_COLS = {"RTO %", "DTO %", "Return %"}
INT_COLS = {"Total Order", "Delivered", "RTO", "Customer Return", "Exchange",
            "Cancelled", "Recovery Qty", "Claim Qty", "P.Count"}


def _write_sheet(writer, name: str, df: pd.DataFrame, total_row: bool = True):
    if df is None or df.empty:
        return
    df.to_excel(writer, sheet_name=name[:31], index=False, startrow=1)
    ws = writer.sheets[name[:31]]

    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=2, column=j)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = max(len(str(col)) + 2, 11)
        ws.column_dimensions[get_column_letter(j)].width = min(width, 34)

        fmt = PCT_FMT if col in PCT_COLS else INT_FMT if col in INT_COLS else MONEY_FMT
        if pd.api.types.is_numeric_dtype(df[col]):
            for i in range(len(df)):
                c = ws.cell(row=3 + i, column=j)
                c.number_format, c.font, c.border = fmt, BODY_FONT, BORDER
        else:
            for i in range(len(df)):
                c = ws.cell(row=3 + i, column=j)
                c.font, c.border = BODY_FONT, BORDER

    if total_row and len(df):
        r = len(df) + 3
        ws.cell(row=r, column=1, value="TOTAL").font = TOTAL_FONT
        ws.cell(row=r, column=1).fill = TOTAL_FILL
        for j, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=r, column=j)
            cell.fill, cell.font = TOTAL_FILL, TOTAL_FONT
            if j > 1 and pd.api.types.is_numeric_dtype(df[col]) and col not in PCT_COLS:
                letter = get_column_letter(j)
                cell.value = f"=SUM({letter}3:{letter}{len(df) + 2})"
                cell.number_format = INT_FMT if col in INT_COLS else MONEY_FMT

    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(df.columns))}{len(df) + 2}"


def write_report(path: Path, *, kpis: dict, sku: pd.DataFrame, state: pd.DataFrame,
                 suborder: pd.DataFrame | None = None,
                 overcharge: pd.DataFrame | None = None,
                 pending: pd.DataFrame | None = None,
                 exceptions: pd.DataFrame | None = None,
                 validation: pd.DataFrame | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        kdf = pd.DataFrame({"KPI": list(kpis), "Value": list(kpis.values())})
        _write_sheet(writer, "Summary", kdf, total_row=False)
        ws = writer.sheets["Summary"]
        for i, k in enumerate(kpis):
            c = ws.cell(row=3 + i, column=2)
            c.number_format = PCT_FMT if "%" in k else MONEY_FMT

        _write_sheet(writer, "SKU Report", sku)
        _write_sheet(writer, "State Report", state)
        if suborder is not None:
            _write_sheet(writer, "Sub Order Report", suborder)
        if overcharge is not None:
            _write_sheet(writer, "Extra Charge", overcharge)
        if pending is not None:
            _write_sheet(writer, "Pending Payments", pending, total_row=False)
        if exceptions is not None:
            _write_sheet(writer, "Exceptions", exceptions, total_row=False)
        if validation is not None:
            _write_sheet(writer, "Validation", validation, total_row=False)
    return path
